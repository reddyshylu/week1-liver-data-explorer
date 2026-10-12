"""Check shared filters and empty states through the Streamlit interface."""
import unittest
from pathlib import Path
from unittest.mock import patch
from streamlit.testing.v1 import AppTest
from liver_data import read_liver

ROOT = Path(__file__).resolve().parents[1]


class LiverAppTests(unittest.TestCase):
    def test_missing_review_download_and_filter_independence(self):
        import io
        from openpyxl import load_workbook
        with patch('streamlit.download_button') as download:
            app = self.app()
            self.assertFalse(app.exception)
            self.assertEqual(app.multiselect[0].value, ['ALT', 'AST'])
            app.sidebar.multiselect[0].set_value([]).run()
            self.assertFalse(app.exception)
            args = [call for call in download.call_args_list
                    if call.args[0] == 'Download Missing Data Review (Excel)'][-1]
            self.assertEqual(args.args[2], 'missing_data_review.xlsx')
            workbook = load_workbook(io.BytesIO(args.args[1]))
            self.assertEqual(workbook['Records to Review'].max_row, 2)
            self.assertEqual(workbook['Summary']['B2'].value, 1)
            app.multiselect[0].set_value([]).run()
            self.assertFalse(app.exception)
            self.assertTrue(any('Select one or more columns' in item.value for item in app.info))

    def app(self):
        return AppTest.from_file(str(ROOT / 'liver_app.py'), default_timeout=30).run()

    def test_combined_filters_and_empty_selections(self):
        app = self.app()
        self.assertFalse(app.exception)
        self.assertEqual(app.metric[0].value, '615')
        app.sidebar.multiselect[0].set_value(['Blood Donor'])
        app.sidebar.multiselect[1].set_value(['Female'])
        app.sidebar.slider[0].set_value((30.0, 40.0)).run()
        self.assertFalse(app.exception)
        frame = read_liver(ROOT / 'data' / 'hcvdat0.csv')
        expected = frame[(frame.Category == 'Blood Donor') & (frame.Sex == 'Female') & frame.Age.between(30, 40)]
        self.assertEqual(int(app.metric[0].value), len(expected))
        self.assertEqual(int(app.metric[2].value), int(expected.isna().sum().sum()))
        summary = app.dataframe[-1].value
        self.assertEqual(int(summary['Count'].sum()), int(expected.ALT.notna().sum()))
        for index in [0, 1]:
            app.sidebar.multiselect[index].set_value([]).run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.metric), 0)
            self.assertTrue(any('No records match' in item.value for item in app.info))
            self.assertTrue(app.get('download_button')[-1].proto.disabled)
            app.sidebar.multiselect[index].set_value(['Blood Donor'] if index == 0 else ['Female']).run()
        app.sidebar.slider[0].set_value((19.0, 19.0)).run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.metric), 0)
        self.assertTrue(app.get('download_button')[-1].proto.disabled)

    def test_export_matches_shared_filter(self):
        app = AppTest.from_file(str(ROOT / 'liver_app.py'), default_timeout=30)
        with patch('streamlit.delta_generator.DeltaGenerator.download_button') as download:
            app.run()
            app.sidebar.slider[0].set_value((30.0, 40.0)).run()
            args = download.call_args_list[-1]
            import io
            import pandas as pd
            exported = pd.read_csv(io.BytesIO(args.args[1]))
            self.assertTrue(exported.Age.between(30, 40).all())
            self.assertEqual(len(exported), int(app.metric[0].value))
            self.assertNotIn('Unnamed: 0', exported.columns)

    def test_single_and_missing_ages(self):
        frame = read_liver(ROOT / 'data' / 'hcvdat0.csv').head(3)
        for age in [32.0, float('nan')]:
            with self.subTest(age=age):
                frame['Age'] = age
                with patch('liver_data.read_liver', return_value=frame):
                    app = self.app()
                    self.assertFalse(app.exception)
                    self.assertEqual(app.metric[0].value, '3')
                    self.assertEqual(len(app.sidebar.slider), 0)
                    if frame.Age.isna().all():
                        app.sidebar.checkbox[0].set_value(False).run()
                        self.assertFalse(app.exception)
                        self.assertEqual(len(app.metric), 0)


if __name__ == '__main__':
    unittest.main()
