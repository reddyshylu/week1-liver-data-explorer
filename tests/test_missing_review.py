import io
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from liver_data import read_liver, missing_data_review, missing_review_workbook


class MissingReviewTests(unittest.TestCase):
    def test_tokens_overlap_source_preservation_and_workbook(self):
        source = read_liver(Path(__file__).resolve().parents[1] / 'data/hcvdat0.csv').head(7)
        source['ID'] = ['001', '002', '003', '004', '005', '006', '007']
        source['Note'] = '=1+1'
        source['ALT'] = ['', '   ', 'NA', 'N/A', 'null', 0, 1]
        source['AST'] = ['NA', 2, 3, 4, 5, 0, ' null ']
        csv = source.to_csv(index=False)
        frame = read_liver(io.StringIO(csv))
        before = frame.copy(deep=True)
        records, summary = missing_data_review(frame, ['ALT', 'AST'])
        self.assertEqual(summary['Missing count'].tolist(), [5, 2])
        self.assertEqual(summary['Missing percentage'].tolist(), [5 / 7, 2 / 7])
        self.assertEqual(records['CSV row number'].tolist(), [2, 3, 4, 5, 6, 8])
        self.assertEqual(records['Missing fields'].tolist(), ['ALT, AST', 'ALT', 'ALT', 'ALT', 'ALT', 'AST'])
        self.assertEqual(len(records), 6)
        pd.testing.assert_frame_equal(frame, before)
        workbook = load_workbook(io.BytesIO(missing_review_workbook(records, summary)))
        self.assertEqual(workbook.sheetnames, ['Records to Review', 'Summary'])
        sheet = workbook['Records to Review']
        self.assertEqual(sheet.max_row, 7)
        headers = [cell.value for cell in sheet[1]]
        self.assertIn('ID', headers)
        self.assertEqual(sheet.cell(2, headers.index('ID') + 1).value, '001')
        self.assertEqual(sheet.cell(2, headers.index('Note') + 1).data_type, 's')
        self.assertEqual(workbook['Summary']['B2'].value, 5)
        self.assertAlmostEqual(workbook['Summary']['C2'].value, 5 / 7)
        self.assertEqual(workbook['Summary']['C2'].number_format, '0.0%')

    def test_empty_and_no_missing_selection(self):
        frame = pd.DataFrame({'ALT': [0, 1], 'AST': [0, 2]})
        for columns in [[], ['ALT'], ['ALT', 'AST']]:
            records, summary = missing_data_review(frame, columns)
            self.assertTrue(records.empty)
            workbook = load_workbook(io.BytesIO(missing_review_workbook(records, summary)))
            self.assertEqual(workbook['Records to Review'].max_row, 1)

    def test_full_dataset_includes_late_rows(self):
        frame = read_liver(Path(__file__).resolve().parents[1] / 'data/hcvdat0.csv')
        records, summary = missing_data_review(frame, ['ALT', 'AST'])
        self.assertEqual(summary['Missing count'].tolist(), [1, 0])
        self.assertEqual(len(records), 1)
        self.assertTrue((records['CSV row number'] > 51).all())


if __name__ == '__main__':
    unittest.main()
