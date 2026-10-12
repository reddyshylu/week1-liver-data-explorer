"""CSV validation for the Liver Data Explorer."""
import numpy as np
import pandas as pd

NUMERIC = ['Age', 'ALB', 'ALP', 'ALT', 'AST', 'BIL', 'CHE', 'CHOL', 'CREA', 'GGT', 'PROT']
REQUIRED = ['Category', 'Sex', *NUMERIC]


def read_liver(source):
    """Read a liver CSV, preserving missing measurements without imputation."""
    try:
        frame = pd.read_csv(source, dtype=str, keep_default_na=False, skip_blank_lines=False)
    except (ValueError, OSError, UnicodeError) as error:
        raise ValueError('Could not read this CSV. Please upload a valid CSV file.') from error
    missing = [name for name in REQUIRED if name not in frame.columns]
    if missing:
        raise ValueError('Missing required columns: ' + ', '.join(missing))
    if frame.empty:
        raise ValueError('The CSV contains no records.')
    source_records = frame.copy()
    frame = frame[REQUIRED].copy()
    for name in frame.columns:
        frame[name] = frame[name].map(lambda value: np.nan if is_missing(value) else value)
    for name in NUMERIC:
        original = frame[name]
        converted = pd.to_numeric(original, errors='coerce')
        if (original.notna() & converted.isna()).any() or np.isinf(converted).any():
            raise ValueError(f'{name} must contain finite numbers or blank values.')
        frame[name] = converted
    for name in ['Category', 'Sex']:
        if frame[name].isna().any() or frame[name].astype(str).str.strip().eq('').any():
            raise ValueError(f'{name} must be present for every record.')
        frame[name] = frame[name].astype(str).str.strip()
    frame['Category'] = frame['Category'].replace({
        '0=Blood Donor': 'Blood Donor', '0s=suspect Blood Donor': 'Suspect Blood Donor',
        '1=Hepatitis': 'Hepatitis', '2=Fibrosis': 'Fibrosis', '3=Cirrhosis': 'Cirrhosis',
    })
    frame['Sex'] = frame['Sex'].replace({'m': 'Male', 'f': 'Female'})
    frame.attrs['source_records'] = source_records.to_dict('list')
    return frame


def is_missing(value):
    """Recognize explicit missing tokens without treating zero as missing."""
    return bool(pd.isna(value)) or (isinstance(value, str) and value.strip().lower() in {'', 'na', 'n/a', 'null'})


def missing_data_review(frame, columns):
    """Review original loaded records, independently of dashboard filters."""
    source = pd.DataFrame(frame.attrs.get('source_records', frame.to_dict('list')))
    columns = list(dict.fromkeys(columns))
    mask = source[columns].apply(lambda series: series.map(is_missing))
    affected = mask.any(axis=1)
    records = source.loc[affected].copy()
    # Avoid overwriting user-supplied columns with report metadata.
    row_label = 'CSV row number'
    while row_label in records.columns:
        row_label = 'Review ' + row_label
    if 'Missing fields' in records.columns:
        original_label = 'Original Missing fields'
        while original_label in records.columns:
            original_label = 'Original ' + original_label
        records = records.rename(columns={'Missing fields': original_label})
    records.insert(0, row_label, source.index[affected] + 2)
    records['Missing fields'] = mask.loc[affected].apply(
        lambda row: ', '.join(name for name in columns if row[name]), axis=1)
    summary = pd.DataFrame({'Column': columns,
                            'Missing count': [int(mask[name].sum()) for name in columns],
                            'Missing percentage': [float(mask[name].mean()) if len(source) else 0.0 for name in columns]})
    return records, summary


def missing_review_workbook(records, summary):
    """Create a separate Excel report in memory; never write the source CSV."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, data in [('Records to Review', records), ('Summary', summary)]:
        sheet = workbook.create_sheet(title)
        sheet.append(list(data.columns))
        for row in data.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in row])
        # CSV text is literal data, including strings beginning with '='.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for index, name in enumerate(data.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(50, max(18, len(str(name)) + 3))
        if title == 'Summary':
            for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
                row[0].number_format = '0.0%'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
